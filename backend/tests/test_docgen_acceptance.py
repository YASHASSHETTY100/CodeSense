"""Tests for DOCX and PDF document generation acceptance."""
from __future__ import annotations
import os, tempfile
import docx
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app.models.models import Base, Project, SourceFile
from app.analysis.structural.pipeline import run_structural_analysis
from app.analysis.functional.pipeline import run_functional_analysis
from app.knowledge_base.pipeline import run_kb_indexing
from app.docgen.docgen import generate_doc, module_facts, build_outline


def _db():
    eng = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=eng)
    return sessionmaker(bind=eng)()


def _setup_repo(db):
    d = tempfile.mkdtemp()
    p = Project(name="DocTestProject", repo_url="https://github.com/test/doctest")
    db.add(p)
    db.flush()

    files = {
        "cases/api.py": """
from fastapi import APIRouter
router = APIRouter()

@router.post("/cases/submit")
def submit_case(case_id: int):
    \"\"\"Submit missing person case.\"\"\"
    if case_id <= 0:
        raise ValueError("Invalid case ID")
    return {"status": "created"}
""",
        "cases/models.py": """
class CaseRecord:
    case_id: int
    name: str
""",
        "requirements.txt": """fastapi==0.115.0\nsqlite3\n""",
        "README.md": """# Case Management System\nCoordinates missing-person records.\n""",
    }

    for rel, content in files.items():
        fp = os.path.join(d, rel)
        os.makedirs(os.path.dirname(fp), exist_ok=True)
        open(fp, "w", encoding="utf-8").write(content)
        lang = "python" if rel.endswith(".py") else "markdown"
        db.add(SourceFile(project_id=p.id, path=rel, language=lang, hash="h"))

    db.commit()
    run_structural_analysis(db, p.id, d)
    run_functional_analysis(db, p.id, d)
    run_kb_indexing(db, p.id, d)
    return p, d


def test_docx_generation_and_structure():
    db = _db()
    p, _ = _setup_repo(db)

    path = generate_doc(db, p.id, p.name, "cases", "docx", 1, 991)
    assert os.path.exists(path)
    assert os.path.getsize(path) > 1000

    doc = docx.Document(path)
    assert len(doc.paragraphs) >= 15
    headings = [p.text for p in doc.paragraphs if p.text.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.", "11.", "12.", "13."))]
    assert len(headings) >= 10
    assert any("Project Purpose" in h for h in headings)
    assert any("Business Overview" in h for h in headings)
    assert any("Workflows" in h for h in headings)


def test_pdf_generation_and_size():
    db = _db()
    p, _ = _setup_repo(db)

    path = generate_doc(db, p.id, p.name, "cases", "pdf", 1, 992)
    assert os.path.exists(path)
    assert os.path.getsize(path) > 1500
